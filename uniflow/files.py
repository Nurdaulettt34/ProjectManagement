"""File validation and text extraction. Works on bytes in memory only:
nothing is written to the local disk."""
import hashlib
import io
import zipfile

ALLOWED_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
MIN_TEXT_CHARS = 20  # less than this is treated as "no readable text"


class UnreadableDocument(Exception):
    pass


def file_extension(filename):
    if not filename or "." not in filename:
        return ""
    return filename.rsplit(".", 1)[1].lower()


def sha256_hex(data):
    return hashlib.sha256(data).hexdigest()


def has_valid_signature(data, ext):
    """Check the file *content*, not just its name (US2).
    PDF: '%PDF-' magic within the first 1 KB (allowed by the PDF spec).
    DOCX: a ZIP archive that really contains word/document.xml."""
    if ext == "pdf":
        return b"%PDF-" in data[:1024]
    if ext == "docx":
        if not data.startswith(b"PK\x03\x04"):
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                return "word/document.xml" in zf.namelist()
        except zipfile.BadZipFile:
            return False
    return False


def extract_text(data, ext):
    """Return the document text, or raise UnreadableDocument."""
    try:
        if ext == "pdf":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
            text = "\n\n".join(pages)
        elif ext == "docx":
            import docx
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:  # university forms often live in tables
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if len(cells) == 2 and cells[0] and cells[1]:
                        parts.append(f"{cells[0]}: {cells[1]}")
                    else:
                        parts.append(" ".join(c for c in cells if c))
            text = "\n".join(parts)
        else:
            raise UnreadableDocument("Unsupported file type.")
    except UnreadableDocument:
        raise
    except Exception as exc:
        raise UnreadableDocument("The file could not be read — it may be damaged or password-protected.") from exc

    text = text.strip()
    if len(text) < MIN_TEXT_CHARS:
        raise UnreadableDocument("No readable text found — the document may be scanned, empty or image-only.")
    return text
